import pytest
from docx import Document

from app.services.prd_parser import (
    PRD_MAX_CHARS,
    parse_prd_file,
    truncate_prd_text,
)


def test_truncate_prd_text():
    long = "a" * (PRD_MAX_CHARS + 100)
    result, truncated = truncate_prd_text(long)
    assert truncated is True
    assert len(result) == PRD_MAX_CHARS


def test_truncate_short_text_not_truncated():
    text = "short prd content"
    result, truncated = truncate_prd_text(text)
    assert truncated is False
    assert result == text


def test_parse_docx_extracts_paragraphs(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Feature: User Authentication")
    doc.add_paragraph("Owner: Prerana Shukla")
    doc.add_paragraph("Milestone: Launch auth by Q3")
    doc.save(str(path))

    with open(path, "rb") as f:
        text = parse_prd_file(f.read(), "sample.docx")
    assert "User Authentication" in text
    assert "Prerana Shukla" in text


def test_parse_legacy_doc_rejected():
    with pytest.raises(ValueError, match="Legacy .doc"):
        parse_prd_file(b"fake", "spec.doc")
