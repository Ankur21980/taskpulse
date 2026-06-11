import io
from typing import Tuple, Union

from docx import Document
from pypdf import PdfReader

PRD_MAX_CHARS = 30_000


def truncate_prd_text(text: str) -> Tuple[str, bool]:
    if len(text) <= PRD_MAX_CHARS:
        return text, False
    return text[:PRD_MAX_CHARS], True


def _parse_pdf(content: bytes) -> str:
    reader = PdfReader(io.BytesIO(content))
    pages = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            pages.append(page_text.strip())
    return "\n\n".join(pages)


def _parse_docx(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def parse_prd_file(content: Union[bytes, str], filename: str) -> str:
    lower = filename.lower()
    if lower.endswith(".doc"):
        raise ValueError("Legacy .doc not supported. Save as .docx or PDF.")
    if lower.endswith(".pdf"):
        raw = _parse_pdf(content if isinstance(content, bytes) else content.encode())
    elif lower.endswith(".docx"):
        raw = _parse_docx(content if isinstance(content, bytes) else content.encode())
    else:
        raise ValueError(f"Unsupported PRD format: {filename}. Use .pdf or .docx.")

    if len(raw.strip()) < 50:
        raise ValueError(
            "Could not extract text from document. Use a text-based PDF or .docx."
        )
    return raw
