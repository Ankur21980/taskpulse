from pathlib import Path

from docx import Document

SAMPLES_DIR = Path(__file__).resolve().parent.parent / "samples"


def main() -> None:
    SAMPLES_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("TaskPulse Q3 PRD — 3-Week Sprint", 0)
    doc.add_paragraph(
        "Sprint window: June 9 – June 30, 2026. All deliverables below must fit within this sprint."
    )
    doc.add_heading("Authentication", 1)
    doc.add_paragraph(
        "Requirement: Implement OAuth2 login API. Owner: Prerana Shukla. Due: June 20. Priority: High."
    )
    doc.add_paragraph(
        "Requirement: Build login and signup UI screens. Owner: Hiren Chafekar. Due: June 25."
    )
    doc.add_heading("Upload & Review", 1)
    doc.add_paragraph(
        "Requirement: Refactor upload page into tabbed layout. Owner: Anisha Kumari. Due: June 18."
    )
    doc.add_paragraph(
        "Deliverable: Add unit tests for upload components. Owner: Gowtham L. Due: June 28."
    )
    doc.add_heading("Notifications", 1)
    doc.add_paragraph(
        "Milestone: Email notification service scaffold. Owner: Prerana Shukla. Due: June 30."
    )
    out = SAMPLES_DIR / "sample-prd.docx"
    doc.save(str(out))
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
